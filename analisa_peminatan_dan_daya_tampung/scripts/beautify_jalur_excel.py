import os
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

script_dir = os.path.dirname(os.path.abspath(__file__))
base_data_dir = os.path.abspath(os.path.join(script_dir, '..', 'data'))
excel_master = os.path.join(base_data_dir, 'master_analisa_peminatan_dan_daya_tampung_2022_2026.xlsx')
out_excel = os.path.join(base_data_dir, 'analisa_jalur_masuk_dan_kebocoran_per_prodi_2022_2026.xlsx')

print("Loading data from master Excel...")
df_all = pd.read_excel(excel_master, sheet_name='Rincian_Jalur_Semua_Tahun')
years = [2022, 2023, 2024, 2025, 2026]
jalurs = ['SNBP', 'SNBT', 'SMMPTN', 'TALENTA', 'SMC', 'ADIK']

wb = openpyxl.Workbook()
# remove default sheet
wb.remove(wb.active)

# Color styles
NAVY_HEADER = "1E3A8A"
SLATE_SUB = "334155"
ICE_ROW = "F8FAFC"
WHITE = "FFFFFF"
BORDER_COLOR = "CBD5E1"
RED_ALERT_FILL = "FEE2E2"
RED_ALERT_FONT = "991B1B"
GREEN_GOOD_FILL = "DCFCE7"
GREEN_GOOD_FONT = "166534"
GOLD_ACCENT = "FEF3C7"
GOLD_FONT = "92400E"

thin_side = Side(style='thin', color=BORDER_COLOR)
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
thick_bottom = Border(left=thin_side, right=thin_side, top=thin_side, bottom=Side(style='medium', color='1E3A8A'))

font_title = Font(name='Segoe UI', size=13, bold=True, color='FFFFFF')
font_sub = Font(name='Segoe UI', size=9.5, italic=True, color='E2E8F0')
font_th = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
font_cell = Font(name='Segoe UI', size=9.5, color='0F172A')
font_bold = Font(name='Segoe UI', size=9.5, bold=True, color='0F172A')

fill_navy = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type='solid')
fill_slate = PatternFill(start_color=SLATE_SUB, end_color=SLATE_SUB, fill_type='solid')
fill_ice = PatternFill(start_color=ICE_ROW, end_color=ICE_ROW, fill_type='solid')
fill_white = PatternFill(start_color=WHITE, end_color=WHITE, fill_type='solid')
fill_alert = PatternFill(start_color=RED_ALERT_FILL, end_color=RED_ALERT_FILL, fill_type='solid')
fill_good = PatternFill(start_color=GREEN_GOOD_FILL, end_color=GREEN_GOOD_FILL, fill_type='solid')
fill_gold = PatternFill(start_color=GOLD_ACCENT, end_color=GOLD_ACCENT, fill_type='solid')

align_center = Alignment(horizontal='center', vertical='center')
align_left = Alignment(horizontal='left', vertical='center')
align_right = Alignment(horizontal='right', vertical='center')
align_wrap_center = Alignment(horizontal='center', vertical='center', wrap_text=True)

def apply_title_banner(ws, title_text, subtitle_text, num_cols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols)
    
    cell_t = ws.cell(row=1, column=1, value=title_text)
    cell_t.font = font_title
    cell_t.fill = fill_navy
    cell_t.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    
    cell_s = ws.cell(row=2, column=1, value=subtitle_text)
    cell_s.font = font_sub
    cell_s.fill = fill_navy
    cell_s.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 10

def auto_fit_columns(ws, max_cols=None, min_width=11):
    cols_to_check = range(1, max_cols + 1) if max_cols else range(1, ws.max_column + 1)
    for col in cols_to_check:
        col_letter = get_column_letter(col)
        max_len = 0
        for row in range(4, ws.max_row + 1):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                s_val = str(val)
                if len(s_val) > max_len:
                    max_len = len(s_val)
        hdr_val = str(ws.cell(row=4, column=col).value or '')
        hdr_len = max(len(w) for w in hdr_val.split('\n')) if hdr_val else 0
        needed_width = max(max_len + 3, hdr_len + 4, min_width)
        ws.column_dimensions[col_letter].width = min(needed_width, 42)

# =============================================================
# SHEET 1: RINGKASAN MAKRO JALUR
# =============================================================
print("Building Styled Sheet 1: Ringkasan_Makro_Jalur...")
ws1 = wb.create_sheet(title="Ringkasan_Makro_Jalur")
ws1.sheet_properties.tabColor = "F59E0B" # Amber

rekap_makro = []
for yr in years:
    df_yr = df_all[df_all['Tahun Akademik'] == yr]
    for j in jalurs:
        sub = df_yr[df_yr['Jalur Penerimaan'] == j]
        if len(sub) > 0:
            pem = int(sub['Jumlah Peminat'].sum())
            dt = int(sub['Target Daya Tampung'].sum())
            lulus = int(sub['Calon Lulus Seleksi'].sum())
            du = int(sub['Mahasiswa Daftar Ulang'].sum())
            gugur = int(sub['Mundur / Gugur'].sum())
            y_rate = du / lulus if lulus > 0 else 0
            d_rate = gugur / lulus if lulus > 0 else 0
            rekap_makro.append([yr, j, pem, dt, lulus, du, gugur, y_rate, d_rate])

# 5-year totals
for j in jalurs:
    sub = df_all[df_all['Jalur Penerimaan'] == j]
    if len(sub) > 0:
        pem = int(sub['Jumlah Peminat'].sum())
        dt = int(sub['Target Daya Tampung'].sum())
        lulus = int(sub['Calon Lulus Seleksi'].sum())
        du = int(sub['Mahasiswa Daftar Ulang'].sum())
        gugur = int(sub['Mundur / Gugur'].sum())
        y_rate = du / lulus if lulus > 0 else 0
        d_rate = gugur / lulus if lulus > 0 else 0
        rekap_makro.append(["Total 5-Tahun", j, pem, dt, lulus, du, gugur, y_rate, d_rate])

headers_s1 = [
    "Tahun Akademik", "Jalur Penerimaan", "Jumlah Peminat\n(Pelamar)", 
    "Target Kuota\n(Daya Tampung)", "Calon Lulus\nSeleksi", "Mahasiswa Masuk\n(Daftar Ulang)", 
    "Calon Gugur\n(Mundur)", "Yield Rate\n(Tingkat Konversi)", "Tingkat Kebocoran\n(Drop-out Rate)"
]

apply_title_banner(ws1, 
                   "UNIVERSITAS SYIAH KUALA — REKAPITULASI MAKRO JALUR PENERIMAAN (2022–2026)",
                   "Dekomposisi Volume Pelamar, Kuota, Kelulusan, dan Evaluasi Efisiensi Konversi per Pintu Masuk",
                   len(headers_s1))

# Write Headers at row 4
ws1.row_dimensions[4].height = 28
for col_idx, h in enumerate(headers_s1, start=1):
    c = ws1.cell(row=4, column=col_idx, value=h)
    c.font = font_th
    c.fill = fill_slate
    c.alignment = align_wrap_center
    c.border = thick_bottom

# Write Data rows starting row 5
for row_idx, r_data in enumerate(rekap_makro, start=5):
    ws1.row_dimensions[row_idx].height = 20
    is_total_5y = str(r_data[0]).startswith("Total")
    row_fill = fill_gold if is_total_5y else (fill_ice if row_idx % 2 == 1 else fill_white)
    
    for c_idx, val in enumerate(r_data, start=1):
        c = ws1.cell(row=row_idx, column=c_idx, value=val)
        c.border = thin_border
        c.fill = row_fill
        c.font = font_bold if is_total_5y else font_cell
        
        # Formatting
        if c_idx in [1, 2]:
            c.alignment = align_center
        elif c_idx in [3, 4, 5, 6, 7]:
            c.alignment = align_right
            c.number_format = '#,##0'
        elif c_idx in [8, 9]:
            c.alignment = align_right
            c.number_format = '0.0%'
            # Conditional styling
            if not is_total_5y:
                if c_idx == 8 and val >= 0.90:
                    c.fill = fill_good
                    c.font = Font(name='Segoe UI', size=9.5, bold=True, color=GREEN_GOOD_FONT)
                elif c_idx == 9 and val >= 0.35:
                    c.fill = fill_alert
                    c.font = Font(name='Segoe UI', size=9.5, bold=True, color=RED_ALERT_FONT)

auto_fit_columns(ws1, len(headers_s1))
ws1.freeze_panes = 'C5'
ws1.auto_filter.ref = f"A4:{get_column_letter(len(headers_s1))}{ws1.max_row}"

# =============================================================
# SHEET 2: PIVOT 5 TAHUN PER PROGRAM STUDI
# =============================================================
print("Building Styled Sheet 2: Pivot_5Thn_Per_Prodi...")
ws2 = wb.create_sheet(title="Pivot_5Thn_Per_Prodi")
ws2.sheet_properties.tabColor = "1E3A8A" # Deep Navy

prodi_list = df_all[['Fakultas', 'Nama Program Studi', 'Jenjang']].drop_duplicates().sort_values(by=['Fakultas', 'Nama Program Studi'])

headers_s2 = [
    "No", "Fakultas", "Nama Program Studi", "Jenjang",
    "DU SNBP", "DU SNBT", "DU Mandiri", "DU Talenta", "DU SMC", "DU ADIK", "Total DU (5-Thn)",
    "Gugur SNBP", "Gugur SNBT", "Gugur Mandiri", "Gugur Talenta", "Gugur SMC", "Gugur ADIK", "Total Gugur (5-Thn)",
    "Total Lulus", "Yield Rate (5-Thn)", "Tingkat Kebocoran", "Pintu Masuk Dominan"
]

apply_title_banner(ws2,
                   "MATRIKS 5-TAHUNAN PENERIMAAN & KEBOCORAN PER PROGRAM STUDI (2022–2026)",
                   "Distribusi Mahasiswa Masuk vs Calon Mangkir per Jalur Penerimaan untuk Seluruh Program Studi USK",
                   len(headers_s2))

# Multi-level header row 4
ws2.row_dimensions[4].height = 28
for col_idx, h in enumerate(headers_s2, start=1):
    c = ws2.cell(row=4, column=col_idx, value=h)
    c.font = font_th
    c.alignment = align_wrap_center
    c.border = thick_bottom
    if "DU" in h or "Total DU" in h:
        c.fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type='solid') # Blue
    elif "Gugur" in h or "Total Gugur" in h:
        c.fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type='solid') # Red
    elif "Yield" in h or "Kebocoran" in h:
        c.fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type='solid') # Teal
    else:
        c.fill = fill_slate

row_num = 5
for p_idx, (_, p) in enumerate(prodi_list.iterrows(), start=1):
    fak = p['Fakultas']
    nm = p['Nama Program Studi']
    jnj = p['Jenjang']
    
    sub = df_all[(df_all['Fakultas'] == fak) & (df_all['Nama Program Studi'] == nm) & (df_all['Jenjang'] == jnj)]
    
    du_vals = {}
    gugur_vals = {}
    total_lulus = int(sub['Calon Lulus Seleksi'].sum())
    total_du = int(sub['Mahasiswa Daftar Ulang'].sum())
    total_gugur = int(sub['Mundur / Gugur'].sum())
    
    for j in jalurs:
        sub_j = sub[sub['Jalur Penerimaan'] == j]
        du_vals[j] = int(sub_j['Mahasiswa Daftar Ulang'].sum())
        gugur_vals[j] = int(sub_j['Mundur / Gugur'].sum())
        
    y_rate = total_du / total_lulus if total_lulus > 0 else 0
    d_rate = total_gugur / total_lulus if total_lulus > 0 else 0
    
    if total_du > 0:
        max_j = max(du_vals, key=du_vals.get)
        pct_dom = round(du_vals[max_j] / total_du * 100, 1)
        dom_txt = f"{max_j} ({pct_dom}%)"
    else:
        dom_txt = "-"
        
    row_data = [
        p_idx, fak, nm, jnj,
        du_vals['SNBP'], du_vals['SNBT'], du_vals['SMMPTN'], du_vals['TALENTA'], du_vals['SMC'], du_vals['ADIK'], total_du,
        gugur_vals['SNBP'], gugur_vals['SNBT'], gugur_vals['SMMPTN'], gugur_vals['TALENTA'], gugur_vals['SMC'], gugur_vals['ADIK'], total_gugur,
        total_lulus, y_rate, d_rate, dom_txt
    ]
    
    ws2.row_dimensions[row_num].height = 19
    r_fill = fill_ice if row_num % 2 == 1 else fill_white
    
    for c_idx, val in enumerate(row_data, start=1):
        c = ws2.cell(row=row_num, column=c_idx, value=val)
        c.border = thin_border
        c.fill = r_fill
        c.font = font_cell
        
        if c_idx in [1, 4, 22]:
            c.alignment = align_center
        elif c_idx in [2, 3]:
            c.alignment = align_left
        elif c_idx in [20, 21]:
            c.alignment = align_right
            c.number_format = '0.0%'
            if c_idx == 20 and val >= 0.90:
                c.fill = fill_good
                c.font = Font(name='Segoe UI', size=9.5, bold=True, color=GREEN_GOOD_FONT)
            elif c_idx == 21 and val >= 0.35:
                c.fill = fill_alert
                c.font = Font(name='Segoe UI', size=9.5, bold=True, color=RED_ALERT_FONT)
        else:
            c.alignment = align_right
            c.number_format = '#,##0'
            if c_idx == 11: # Total DU
                c.font = font_bold
                c.fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type='solid')
            elif c_idx == 18: # Total Gugur
                c.font = font_bold
                c.fill = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type='solid')
    row_num += 1

auto_fit_columns(ws2, len(headers_s2))
ws2.freeze_panes = 'E5'
ws2.auto_filter.ref = f"A4:{get_column_letter(len(headers_s2))}{ws2.max_row}"

# =============================================================
# SHEET 3: DETAIL 5 TAHUN PER PRODI X JALUR
# =============================================================
print("Building Styled Sheet 3: Detail_5Thn_Prodi_Jalur...")
ws3 = wb.create_sheet(title="Detail_5Thn_Prodi_Jalur")
ws3.sheet_properties.tabColor = "3B82F6" # Blue

headers_s3 = [
    "No", "Fakultas", "Nama Program Studi", "Jenjang", "Jalur Penerimaan",
    "Jumlah Peminat", "Target Kuota (DT)", "Calon Lulus Seleksi", "Mahasiswa Daftar Ulang",
    "Calon Gugur (Mundur)", "Yield Rate (%)", "Tingkat Kebocoran (%)"
]

apply_title_banner(ws3,
                   "DETAIL TABULAR AKUMULASI 5 TAHUN PER PROGRAM STUDI & JALUR (2022–2026)",
                   "Pembedahan Menyeluruh Metrik Seleksi dan Tingkat Konversi Tiap Pintu Masuk di Tingkat Program Studi",
                   len(headers_s3))

ws3.row_dimensions[4].height = 28
for col_idx, h in enumerate(headers_s3, start=1):
    c = ws3.cell(row=4, column=col_idx, value=h)
    c.font = font_th
    c.fill = fill_slate
    c.alignment = align_wrap_center
    c.border = thick_bottom

df_5y_det = df_all.groupby(['Fakultas', 'Nama Program Studi', 'Jenjang', 'Jalur Penerimaan']).agg({
    'Jumlah Peminat': 'sum',
    'Target Daya Tampung': 'sum',
    'Calon Lulus Seleksi': 'sum',
    'Mahasiswa Daftar Ulang': 'sum',
    'Mundur / Gugur': 'sum'
}).reset_index().sort_values(by=['Fakultas', 'Nama Program Studi', 'Jalur Penerimaan'])

row_num = 5
for idx, (_, r) in enumerate(df_5y_det.iterrows(), start=1):
    pem = int(r['Jumlah Peminat'])
    dt = int(r['Target Daya Tampung'])
    lulus = int(r['Calon Lulus Seleksi'])
    du = int(r['Mahasiswa Daftar Ulang'])
    gugur = int(r['Mundur / Gugur'])
    y_rate = du / lulus if lulus > 0 else 0
    d_rate = gugur / lulus if lulus > 0 else 0
    
    r_data = [idx, r['Fakultas'], r['Nama Program Studi'], r['Jenjang'], r['Jalur Penerimaan'],
              pem, dt, lulus, du, gugur, y_rate, d_rate]
    
    ws3.row_dimensions[row_num].height = 19
    r_fill = fill_ice if row_num % 2 == 1 else fill_white
    
    for c_idx, val in enumerate(r_data, start=1):
        c = ws3.cell(row=row_num, column=c_idx, value=val)
        c.border = thin_border
        c.fill = r_fill
        c.font = font_cell
        
        if c_idx in [1, 4, 5]:
            c.alignment = align_center
        elif c_idx in [2, 3]:
            c.alignment = align_left
        elif c_idx in [11, 12]:
            c.alignment = align_right
            c.number_format = '0.0%'
            if c_idx == 11 and val >= 0.90:
                c.fill = fill_good
                c.font = Font(name='Segoe UI', size=9.5, bold=True, color=GREEN_GOOD_FONT)
            elif c_idx == 12 and val >= 0.35:
                c.fill = fill_alert
                c.font = Font(name='Segoe UI', size=9.5, bold=True, color=RED_ALERT_FONT)
        else:
            c.alignment = align_right
            c.number_format = '#,##0'
    row_num += 1

auto_fit_columns(ws3, len(headers_s3))
ws3.freeze_panes = 'F5'
ws3.auto_filter.ref = f"A4:{get_column_letter(len(headers_s3))}{ws3.max_row}"

# =============================================================
# SHEETS 4-8: RINCIAN TAHUNAN (2026 s.d. 2022)
# =============================================================
annual_colors = {
    2026: "10B981", # Emerald
    2025: "64748B",
    2024: "64748B",
    2023: "64748B",
    2022: "64748B"
}

headers_annual = [
    "No", "Tahun", "Fakultas", "Nama Program Studi", "Jenjang", "Jalur Masuk",
    "Jumlah Peminat", "Target Kuota (DT)", "Calon Lulus Seleksi", "Mahasiswa Daftar Ulang",
    "Calon Gugur (Mundur)", "Yield Rate (%)", "Tingkat Kebocoran (%)"
]

for yr in [2026, 2025, 2024, 2023, 2022]:
    sheet_title = f"Rincian_Tahun_{yr}"
    print(f"Building Styled Sheet: {sheet_title}...")
    ws_yr = wb.create_sheet(title=sheet_title)
    ws_yr.sheet_properties.tabColor = annual_colors[yr]
    
    apply_title_banner(ws_yr,
                       f"DATA PENERIMAAN & KEBOCORAN MAHASISWA BARU USK — TAHUN AKADEMIK {yr}",
                       f"Rincian Pelamar, Kuota, Kelulusan, dan Mahasiswa Masuk per Program Studi x Jalur Masuk Tahun {yr}",
                       len(headers_annual))
    
    ws_yr.row_dimensions[4].height = 28
    for col_idx, h in enumerate(headers_annual, start=1):
        c = ws_yr.cell(row=4, column=col_idx, value=h)
        c.font = font_th
        c.fill = fill_slate
        c.alignment = align_wrap_center
        c.border = thick_bottom
        
    df_yr = df_all[df_all['Tahun Akademik'] == yr].sort_values(by=['Fakultas', 'Nama Program Studi', 'Jalur Penerimaan'])
    
    row_num = 5
    for idx, (_, r) in enumerate(df_yr.iterrows(), start=1):
        pem = int(r['Jumlah Peminat'])
        dt = int(r['Target Daya Tampung'])
        lulus = int(r['Calon Lulus Seleksi'])
        du = int(r['Mahasiswa Daftar Ulang'])
        gugur = int(r['Mundur / Gugur'])
        y_rate = du / lulus if lulus > 0 else 0
        d_rate = gugur / lulus if lulus > 0 else 0
        
        r_data = [idx, yr, r['Fakultas'], r['Nama Program Studi'], r['Jenjang'], r['Jalur Penerimaan'],
                  pem, dt, lulus, du, gugur, y_rate, d_rate]
        
        ws_yr.row_dimensions[row_num].height = 19
        r_fill = fill_ice if row_num % 2 == 1 else fill_white
        
        for c_idx, val in enumerate(r_data, start=1):
            c = ws_yr.cell(row=row_num, column=c_idx, value=val)
            c.border = thin_border
            c.fill = r_fill
            c.font = font_cell
            
            if c_idx in [1, 2, 5, 6]:
                c.alignment = align_center
            elif c_idx in [3, 4]:
                c.alignment = align_left
            elif c_idx in [12, 13]:
                c.alignment = align_right
                c.number_format = '0.0%'
                if c_idx == 12 and val >= 0.90:
                    c.fill = fill_good
                    c.font = Font(name='Segoe UI', size=9.5, bold=True, color=GREEN_GOOD_FONT)
                elif c_idx == 13 and val >= 0.35:
                    c.fill = fill_alert
                    c.font = Font(name='Segoe UI', size=9.5, bold=True, color=RED_ALERT_FONT)
            else:
                c.alignment = align_right
                c.number_format = '#,##0'
        row_num += 1
        
    auto_fit_columns(ws_yr, len(headers_annual))
    ws_yr.freeze_panes = 'G5'
    ws_yr.auto_filter.ref = f"A4:{get_column_letter(len(headers_annual))}{ws_yr.max_row}"

wb.save(out_excel)
print("Magnificent Styled Excel successfully saved at:", out_excel)
